"""Atomicity gates for generated manuscript, table and benchmark-figure artifacts."""
from __future__ import annotations

import importlib.util
import re
import sys
from collections import Counter
from pathlib import Path
from xml.etree import ElementTree

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pytest
from docx import Document
from docx.oxml.ns import qn
from PySide6.QtCore import QByteArray
from PySide6.QtGui import QColor, QImage, QPainter
from PySide6.QtSvg import QSvgRenderer

from teagle_core.artifact_metadata import media_metadata_violations, ooxml_metadata_violations


ROOT = Path(__file__).resolve().parents[1]
_URL_REF = re.compile(r"url\(\s*(['\"]?)#([^)'\"\s]+)\1\s*\)", re.I)
_CSS_REF = re.compile(r"#((?:\\.|[A-Za-z0-9_-])+)(?=\s*[,\{])")


def _load(relpath: str, name: str):
    spec = importlib.util.spec_from_file_location(name, ROOT / relpath)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


def _svg_contract(payload: bytes):
    root = ElementTree.fromstring(payload)
    ids = []
    references = Counter()
    semantic = []
    for element in root.iter():
        attributes = {}
        for key, value in element.attrib.items():
            local = key.rsplit("}", 1)[-1].casefold()
            if local == "id":
                ids.append(value)
            if local == "href" and value.startswith("#"):
                references[value[1:]] += 1
            references.update(match.group(2) for match in _URL_REF.finditer(value))
            if local in {"aria-labelledby", "aria-describedby"}:
                references.update(value.split())
            elif local in {"begin", "end"}:
                references.update(
                    token.rsplit(".", 1)[0]
                    for token in re.split(r"\s*;\s*", value)
                    if "." in token and not token.rsplit(".", 1)[0].replace(".", "").isdigit()
                )
            attributes[key] = value
        if element.tag.rsplit("}", 1)[-1].casefold() == "style":
            references.update(
                match.group(1).replace("\\.", ".")
                for match in _CSS_REF.finditer(element.text or "")
            )
        semantic.append((element.tag, tuple(sorted(attributes.items())), element.text or ""))
    return ids, references, semantic


def _render_svg(payload: bytes):
    renderer = QSvgRenderer(QByteArray(payload))
    assert renderer.isValid()
    image = QImage(64, 64, QImage.Format.Format_ARGB32)
    image.fill(QColor("white"))
    painter = QPainter(image)
    renderer.render(painter)
    painter.end()
    return tuple(image.pixel(x, y) for y in range(image.height()) for x in range(image.width()))


@pytest.fixture(scope="module")
def figure_writer():
    return _load("benchmarks/figures.py", "teagle_test_benchmark_figures")


@pytest.fixture(scope="module")
def table_writer():
    module = _load("benchmarks/make_tables.py", "teagle_test_make_tables")
    if not hasattr(module, "_write_clean_ooxml"):
        pytest.skip("held benchmark table writer is outside the app release")
    return module


@pytest.fixture(scope="module")
def manuscript_writer():
    if not (ROOT / "manuscript" / "build_docx.py").is_file():
        pytest.skip("gitignored manuscript workspace is absent")
    return _load("manuscript/build_docx.py", "teagle_test_build_docx")


@pytest.mark.parametrize("suffix", [".pdf", ".png", ".svg"])
def test_benchmark_figure_writer_cleans_each_supported_format(figure_writer, tmp_path, suffix):
    fig, ax = plt.subplots(figsize=(1, 1))
    ax.plot([0, 1], [0, 1])
    target = tmp_path / f"figure{suffix}"
    try:
        figure_writer.write_clean_figure(fig, target, dpi=72 if suffix == ".png" else None)
    finally:
        plt.close(fig)
    assert target.is_file()
    assert media_metadata_violations(target) == []
    if suffix == ".svg":
        ids, references, _semantic = _svg_contract(target.read_bytes())
        assert not any(value.startswith("matplotlib.axis_") for value in ids)
        assert set(references).issubset(ids)


@pytest.mark.parametrize("newline", ["\n", "\r\n"])
def test_controlled_matplotlib_svg_preamble_is_the_only_producer_exception(
        figure_writer, tmp_path, newline):
    target = tmp_path / "producer.svg"
    preamble = (
        '<?xml version="1.0" encoding="utf-8" standalone="no"?>' + newline
        + '<!DOCTYPE svg PUBLIC "-//W3C//DTD SVG 1.1//EN"' + newline
        + '  "http://www.w3.org/Graphics/SVG/1.1/DTD/svg11.dtd">' + newline
    ).encode("ascii")
    target.write_bytes(preamble + b'<svg xmlns="http://www.w3.org/2000/svg"/>')
    figure_writer._strip_controlled_matplotlib_svg_preamble(target)
    assert target.read_bytes() == b'<svg xmlns="http://www.w3.org/2000/svg"/>'


def test_unknown_svg_dtd_from_producer_refuses_and_preserves_destination(
        figure_writer, tmp_path, monkeypatch):
    target = tmp_path / "figure.svg"
    target.write_bytes(b"existing figure")
    fig, _ax = plt.subplots(figsize=(1, 1))

    def write_unknown_dtd(path, **_kwargs):
        Path(path).write_bytes(
            b'<?xml version="1.0" encoding="utf-8" standalone="no"?>\n'
            b'<!DOCTYPE svg [<!ENTITY probe "unsafe">]>\n'
            b'<svg xmlns="http://www.w3.org/2000/svg"><text>&probe;</text></svg>'
        )

    monkeypatch.setattr(fig, "savefig", write_unknown_dtd)
    try:
        with pytest.raises(ValueError, match="unknown DTD"):
            figure_writer.write_clean_figure(fig, target)
    finally:
        plt.close(fig)
    assert target.read_bytes() == b"existing figure"
    assert not list(tmp_path.glob(".figure.*.svg"))


def test_axis_id_neutralization_preserves_semantics_pixels_references_and_generic_ids(
        figure_writer, tmp_path):
    path = tmp_path / "axis-ids.svg"
    source = (
        b'<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink" '
        b'width="32" height="32" viewBox="0 0 32 32">'
        b'<defs><clipPath id="matplotlib.axis_1"><rect x="2" y="2" width="28" height="28"/>'
        b'</clipPath><path id="generic.shape" d="M0 0h32v32H0z"/></defs>'
        b'<title id="generic.label">Visible panel</title>'
        b'<style>#matplotlib\\.axis_2 { opacity: 1; }</style>'
        b'<g id="matplotlib.axis_2" clip-path="url(#matplotlib.axis_1)" '
        b'aria-labelledby="generic.label matplotlib.axis_2">'
        b'<use xlink:href="#generic.shape" fill="#336699"/>'
        b'<animate attributeName="opacity" begin="matplotlib.axis_2.click" dur="1s"/>'
        b'</g></svg>'
    )
    path.write_bytes(source)
    before_ids, before_references, before_semantic = _svg_contract(source)
    before_pixels = _render_svg(source)

    figure_writer._neutralize_controlled_matplotlib_svg_ids(path)
    neutral = path.read_bytes()
    neutral_ids, neutral_references, neutral_semantic = _svg_contract(neutral)
    assert {"generic.shape", "generic.label"}.issubset(neutral_ids)
    assert not any(value.startswith("matplotlib.axis_") for value in neutral_ids)
    assert neutral_references == Counter({
        {"matplotlib.axis_1": "plot-axis-1", "matplotlib.axis_2": "plot-axis-2"}.get(key, key): count
        for key, count in before_references.items()
    })
    assert set(neutral_references).issubset(neutral_ids)
    restored_text = neutral.decode("utf-8").replace("#plot-axis-2", "#matplotlib\\.axis_2")
    restored = restored_text.replace("plot-axis-1", "matplotlib.axis_1").replace(
        "plot-axis-2", "matplotlib.axis_2").encode("utf-8")
    assert _svg_contract(restored)[2] == before_semantic
    assert _render_svg(neutral) == before_pixels

    from teagle_core.artifact_metadata import clean_svg_metadata
    clean_svg_metadata(path)
    final_ids, final_references, _final_semantic = _svg_contract(path.read_bytes())
    assert final_ids == neutral_ids
    assert final_references == neutral_references
    assert _render_svg(path.read_bytes()) == before_pixels
    assert b"matplotlib" not in path.read_bytes().lower()
    assert before_ids == ["matplotlib.axis_1", "generic.shape", "generic.label", "matplotlib.axis_2"]


def test_benchmark_figure_validation_failure_preserves_existing_file(figure_writer, tmp_path, monkeypatch):
    target = tmp_path / "figure.pdf"
    target.write_bytes(b"existing figure")
    monkeypatch.setattr(figure_writer, "media_metadata_violations", lambda _path: ["injected defect"])
    fig, ax = plt.subplots(figsize=(1, 1))
    ax.plot([0, 1], [1, 0])
    try:
        with pytest.raises(RuntimeError, match="injected defect"):
            figure_writer.write_clean_figure(fig, target)
    finally:
        plt.close(fig)
    assert target.read_bytes() == b"existing figure"
    assert not list(tmp_path.glob(".figure.*.pdf"))


@pytest.mark.parametrize("kind", ["xlsx", "docx"])
def test_table_validation_failure_preserves_existing_file(table_writer, tmp_path, monkeypatch, kind):
    target = tmp_path / f"table.{kind}"
    target.write_bytes(b"existing table")
    monkeypatch.setattr(table_writer, "ooxml_metadata_violations", lambda _path: ["injected defect"])
    with pytest.raises(RuntimeError, match="injected defect"):
        if kind == "xlsx":
            table_writer.write_xlsx(target, ["Header"], [["value"]], "Table")
        else:
            table_writer.write_docx(target, ["Header"], [["value"]], "Caption")
    assert target.read_bytes() == b"existing table"
    assert not list(tmp_path.glob(f".table.*.{kind}"))


def test_manuscript_validation_failure_preserves_existing_file(
        manuscript_writer, tmp_path, monkeypatch):
    target = tmp_path / "manuscript.docx"
    target.write_bytes(b"existing manuscript")
    monkeypatch.setattr(manuscript_writer, "ooxml_metadata_violations", lambda _path: ["injected defect"])
    doc = Document()
    doc.add_paragraph("Synthetic manuscript content")
    with pytest.raises(RuntimeError, match="injected defect"):
        manuscript_writer._save_clean_docx(doc, target)
    assert target.read_bytes() == b"existing manuscript"
    assert not list(tmp_path.glob(".manuscript.*.docx"))


def test_manuscript_writer_produces_a_verified_package(manuscript_writer, tmp_path):
    target = tmp_path / "manuscript.docx"
    doc = Document()
    doc.add_paragraph("Synthetic manuscript content")
    manuscript_writer._save_clean_docx(doc, target)
    assert ooxml_metadata_violations(target) == []


def test_expanded_table_docx_uses_readable_two_column_panels(table_writer, tmp_path):
    target = tmp_path / "table-s1.docx"
    header = ["Tool", "Version", "Per-call confidence reported"]
    rows = [
        ["RepeatModeler2", "2.0.7", "Three independent axes with a long readable explanation."],
        ["HelitronScanner", "1.0", "A second complete source row."],
    ]
    table_writer.write_expanded_docx(target, header, rows, "Table S1. Synthetic layout gate.")
    document = Document(target)
    assert len(document.tables) == len(rows)
    for table, source_row in zip(document.tables, rows):
        assert len(table.columns) == 2
        assert table.rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader")) is not None
        assert table.rows[0].cells[0].width.mm >= 46
        assert table.rows[0].cells[1].width.mm >= 130
        values = {row.cells[0].text: row.cells[1].text for row in table.rows[1:]}
        assert values == dict(zip(header, source_row))
    assert [p.text for p in document.paragraphs if p.text in {row[0] for row in rows}] == [
        row[0] for row in rows]
    assert ooxml_metadata_violations(target) == []


@pytest.mark.parametrize("name", ["figures.py", "fig_calibration.py", "fig_workflow.py", "make_tables.py"])
def test_reviewer_generator_is_an_exact_copy(name):
    assert (ROOT / "benchmarks" / name).read_bytes() == (
        ROOT / "reviewer_package" / "scripts" / name).read_bytes()


def test_all_benchmark_figure_entrypoints_use_the_atomic_writer():
    for relpath in ("benchmarks/fig_calibration.py", "benchmarks/fig_workflow.py"):
        source = (ROOT / relpath).read_text(encoding="utf-8")
        assert "write_clean_figure(" in source
        assert "fig.savefig(os.path.join(OUT" not in source
