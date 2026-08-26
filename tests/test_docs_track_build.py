"""Documentation claims that must be derived from the build, not restated by hand.

A number written into prose drifts the moment the thing it describes changes: the README claimed a
"21-model Pfam panel" after the panel had grown to 30. These tests fail when a stated figure and the
artefact it describes disagree, so the rot is caught by the suite rather than by a reader.
"""
import csv
import io
import json
import os
import re
import zipfile
from pathlib import Path

import pytest

from teagle_core import domains

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def _readme():
    with io.open(os.path.join(_ROOT, "README.md"), encoding="utf-8") as fh:
        return fh.read()


def _read(relpath):
    path = os.path.join(_ROOT, relpath)
    if not os.path.exists(path) and relpath.replace("\\", "/").startswith(("manuscript/", "report/")):
        pytest.skip(f"{relpath} is an ignored local publication artifact")
    with io.open(path, encoding="utf-8") as fh:
        return fh.read()


def test_public_release_text_omits_internal_review_or_assistant_attribution():
    public_text = "\n".join(_read(path) for path in (
        "CHANGELOG.md", "README.md", "app/README.md", "docs/TRY-IT-YOURSELF.md",
        "report/teagle_manual.tex", "report/teagle_report.tex",
    ))
    prohibited = re.compile(
        r"\b(?:multi-agent|review[ -]loops?|fix[ -]loops?|triple-reviewer|adversarial review|AI-generated|ChatGPT|Codex|Claude|OpenAI)\b",
        re.IGNORECASE,
    )
    assert not prohibited.search(public_text)


def test_public_release_does_not_infer_biological_identity_from_ispcr_priming_sites():
    public_text = "\n".join(_read(path) for path in (
        "CHANGELOG.md", "README.md", "app/README.md", "docs/TRY-IT-YOURSELF.md",
        "report/teagle_manual.tex", "report/teagle_report.tex", "app/native/main.py",
        "app/backend/engine.py", "app/backend/teagle_core/genomepcr.py",
    ))
    assert not re.search(r"\bparalog(?:s|ous)?\b", public_text, re.IGNORECASE)
    for unsupported in ("copy-specific", "family-generic", "locus-specific", "more-diverged copies"):
        assert unsupported not in public_text.lower()
    assert "every genome-wide product is an off-target" not in public_text.lower()


def test_protocol_236_doi_is_labelled_result_informed_in_public_documents():
    for path in ("README.md", "CHANGELOG.md", "report/teagle_manual.tex", "report/teagle_report.tex"):
        text = _read(path).lower()
        assert "10.5281/zenodo.22091013" in text
        assert "result-informed corrective registration" in text


def test_report_reliability_axes_are_teagle_specific_and_not_claimed_independent():
    text = " ".join(_read("report/teagle_report.tex").split())
    changelog = " ".join(_read("CHANGELOG.md").split())
    assert "No transposable-element tool publishes" not in text
    assert "three independent axes" not in text
    assert "three independent, citable" not in text
    assert "missing core module or interrupted reading frame" not in text
    assert "two independent axes" not in changelog
    assert "Reliability: three separate axes" in text


def test_rmblast_dfam_docs_name_consensus_partitions_not_profile_hmm_search():
    report = " ".join(_read("report/teagle_report.tex").split())
    manual = " ".join(_read("report/teagle_manual.tex").split())
    for text in (report, manual):
        assert "dfam40.0.h5" in text
        assert "dfam40.curated.consensus.0.h5" in text
        assert "consensus-sequence" in text
    assert "RMBLAST engine against the installed Dfam library" not in report
    assert "HMM terminology applies to the separate native HMMER/Pfam domain scan" in report


def test_dfam_homology_naming_is_not_described_as_a_classification_layer():
    manual = " ".join(_read("report/teagle_manual.tex").split())
    backend = _read("app/backend/teagle_core/wsl.py")
    assert "classification workflow's family-database layer" not in manual
    assert "family-level TE classification" not in backend
    assert "family-level homology naming" in backend


def test_report_distinguishes_te_family_copy_number_from_an_individual_insertion():
    text = " ".join(_read("report/teagle_report.tex").split())
    assert "A transposable element is rarely a single copy" not in text
    assert "tens to thousands of related copies of each element" not in text
    assert "A TE family may be represented by multiple related genomic copies" in text


def test_report_current_landscape_behavior_is_not_labelled_as_version_381():
    text = " ".join(_read("report/teagle_report.tex").split())
    assert "Version 3.8.1 instead reports" not in text
    assert "\\teagle{} instead reports an independent contig-wise union" in text


def test_report_does_not_require_identical_primer_melting_temperatures():
    text = " ".join(_read("report/teagle_report.tex").split())
    assert "matching the $T_m$ of the two primers is exactly what lets both anneal" not in text
    assert "closely matching their $T_m$ values helps both primers anneal efficiently" in text


def test_accession_version_is_scoped_to_sequence_not_the_entire_record():
    manual = " ".join(_read("report/teagle_manual.tex").split())
    app_readme = " ".join(_read("app/README.md").split())
    assert "fixed, immutable object" not in manual
    assert "explicit accession version" in app_readme
    assert "is immutable" not in app_readme
    for text in (manual, app_readme):
        assert "fixed sequence version" in text


def test_line_endonuclease_uses_ape_terminology():
    public_text = " ".join(_read(path) for path in (
        "CHANGELOG.md", "app/README.md", "report/teagle_manual.tex", "report/teagle_report.tex",
    ))
    assert "apurinic-like" not in public_text
    assert "apurinic/apyrimidinic endonuclease-like (APE)" in public_text


def test_ispcr_candidate_count_is_not_claimed_as_a_conservative_floor():
    public_text = " ".join(_read(path) for path in (
        "CHANGELOG.md", "report/teagle_manual.tex", "report/teagle_report.tex",
        "app/backend/teagle_core/genomepcr.py",
    ))
    assert "conservative floor" not in public_text
    for path in ("report/teagle_manual.tex", "report/teagle_report.tex"):
        assert "neither a validated lower nor upper bound" in " ".join(_read(path).split())


def test_ispcr_specific_behavior_cites_official_ucsc_documentation():
    bibliography = _read("report/references.bib")
    assert "@misc{ucscispcr" in bibliography
    assert "@misc{ucscispcrsource" in bibliography
    assert "@misc{ucscfatotwobit" in bibliography
    assert "https://genome.ucsc.edu/cgi-bin/hgPcr" in bibliography
    assert "https://github.com/ucscGenomeBrowser/kent/blob/49c2b7bc/src/jkOwnLib/gfPcrLib.c" in bibliography
    assert "https://github.com/ucscGenomeBrowser/kent/blob/v482_base/src/utils/faToTwoBit/faToTwoBit.c" in bibliography
    assert "{gfPcrLib}: {PCR} {FASTA} and {BED} output formats" in bibliography
    assert "{faToTwoBit}: {FASTA}-to-2bit converter" in bibliography
    for path in ("report/teagle_manual.tex", "report/teagle_report.tex"):
        text = " ".join(_read(path).split())
        assert "15\\,bp" in text
        assert "\\cite{ucscispcr}" in text
        assert "isPcr's FASTA-header coordinates" in text
        assert "\\cite{ucscispcrsource}" in text
        assert "ucscfatotwobit}" in text


def test_ispcr_acceptance_parameter_is_not_described_as_a_seed_or_count_floor():
    public_text = " ".join(_read(path) for path in (
        "CHANGELOG.md", "report/teagle_manual.tex", "report/teagle_report.tex",
        "app/native/main.py", "app/backend/engine.py", "app/backend/teagle_core/genomepcr.py",
    ))
    assert "perfect-match seed" not in public_text
    assert "perfect-3' seed" not in public_text
    assert "a floor at ≥15 bp" not in public_text
    assert "minimum perfect match" in public_text


def test_bounded_search_docs_do_not_guarantee_host_resource_immunity():
    text = " ".join(_read(path) for path in ("report/teagle_manual.tex", "report/teagle_report.tex"))
    assert "cannot exhaust time or memory" not in text
    assert "cannot exhaust memory" not in text
    assert "search work is bounded by explicit caps" in text


def test_manual_describes_longest_orf_selection_as_a_search_budget_heuristic():
    text = " ".join(_read("report/teagle_manual.tex").split())
    assert "a functional element's coding modules live in its longest ORFs" not in text
    assert "search budget prioritises the twelve longest open reading frames" in text
    assert "excluded shorter ORF may still contain a domain" in text


def test_line_rule_uses_only_line_specific_corroboration():
    for path in ("report/teagle_report.tex", "report/teagle_manual.tex"):
        text = " ".join(_read(path).split())
        assert "RNase~H, gag, protease, or env" not in text
        assert "LINE-specific corroboration" in text
        assert "ORF1p or an accepted same-ORF, N-terminal ORF2p endonuclease--RT arrangement" in text
        assert "R2/R4-type" in text
    assert "R2/R4-type" in _read("README.md")


def _require_manuscript():
    if not os.path.isdir(os.path.join(_ROOT, "manuscript")):
        pytest.skip("manuscript working sources are not distributed in public source checkouts")


def _holdout_rows():
    with io.open(os.path.join(_ROOT, "benchmarks", "corpus_holdout.tsv"), encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh, delimiter="\t"))


def _holdout_scores():
    with io.open(os.path.join(_ROOT, "benchmarks", "raw", "scores_holdout.json"), encoding="utf-8") as fh:
        return json.load(fh)


def _broader_scores():
    with io.open(os.path.join(_ROOT, "benchmarks", "raw", "scores.json"), encoding="utf-8") as fh:
        return json.load(fh)


def _publication_table(name):
    path = os.path.join(_ROOT, "manuscript", "tables", name)
    if not os.path.exists(path):
        pytest.skip("local protocol 2.3.6 publication tables are not present")
    with io.open(path, encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def test_readme_panel_size_matches_the_bundled_profiles():
    stated = re.search(r"(\d+)-model Pfam panel", _readme())
    assert stated, "README no longer states the panel size — update this test if that is intended"
    assert int(stated.group(1)) == len(domains._hmms())


def test_readme_does_not_publish_undeposited_result_values():
    readme = _readme()
    assert "result archive has not received a public DOI" in readme
    for unpublished in (
        "The frozen external holdout contained", "42/43", "41/43", "14/14",
        "Across 1,211 structural truth copies", "Two of 11 controlled off-target scenarios",
    ):
        assert unpublished not in readme


def test_active_comparison_scripts_make_no_preregistration_timing_claim():
    for path in ("benchmarks/head_to_head.py", "reviewer_package/scripts/head_to_head.py"):
        text = _read(path).lower()
        for forbidden in ("pre-registered", "preregistered", "pre-register", "preregister"):
            assert forbidden not in text, f"{path} contains timing claim {forbidden!r}"


# ---------------- report/manual capability enumeration (gitignored working files) ----------------
def _doc(name):
    p = os.path.join(_ROOT, "report", name)
    if not os.path.exists(p):
        return None
    with io.open(p, encoding="utf-8") as fh:
        return fh.read()


@pytest.mark.parametrize("doc", ["teagle_manual.tex", "teagle_report.tex"])
def test_doc_states_the_correct_panel_size(doc):
    """The report and manual both enumerate the panel; the stated model count must equal the bundled set,
    so growing the panel without updating the docs fails the build (README once claimed 21 after it hit 30)."""
    text = _doc(doc)
    if text is None:
        pytest.skip(f"{doc} not present (report/ is a gitignored working dir)")
    n = len(domains.DOMAIN_INFO)
    assert re.search(rf"\b{n}[ ~]?(?:public-domain )?Pfam", text) or re.search(rf"\b{n} Pfam profile", text), \
        f"{doc} does not state the current panel size of {n} Pfam models"


@pytest.mark.parametrize("doc", ["teagle_manual.tex", "teagle_report.tex"])
def test_doc_lists_every_superfamily_the_classifier_can_name(doc):
    """Every superfamily/lineage name classify() can emit must appear in the enumeration, or the docs
    under-state what the tool does. Helitron is deliberately NOT here: its helicase is a detected DOMAIN
    (checked by the accession test) but classify() emits no Helitron superfamily call, so listing it as a
    superfamily would over-state the tool."""
    text = _doc(doc)
    if text is None:
        pytest.skip(f"{doc} not present")
    for name in ("Copia", "Gypsy", "LINE", "DIRS", "Tc1/Mariner", "hAT", "CACTA", "MULE", "IS4"):
        assert name in text, f"{doc} does not name the {name!r} superfamily/group in its enumeration"


@pytest.mark.parametrize("doc", ["teagle_manual.tex", "teagle_report.tex"])
def test_doc_does_not_claim_a_helitron_superfamily_call(doc):
    """classify() detects the Helitron helicase domain but emits no Helitron superfamily call. The docs must
    say so, not imply the tool assigns a Helitron superfamily (guards against re-introducing that overclaim)."""
    import sys as _sys
    _be = os.path.join(_ROOT, "app", "backend")
    if _be not in _sys.path:
        _sys.path.insert(0, _be)
    from teagle_core import classify
    cl = classify.classify([], [{"domain": "HEL", "class": "dna:Helitron",
                                 "nt": [100, 900], "strand": "+", "score": 150.0}])
    assert not cl["te_class"].startswith("HEL"), "classify now emits a Helitron call — update the docs and this test"
    text = _doc(doc)
    if text is None:
        pytest.skip(f"{doc} not present")
    assert "does not currently emit" in text or "no dedicated" in text, \
        f"{doc} must state that no dedicated Helitron superfamily call is emitted"


@pytest.mark.parametrize("doc", ["teagle_manual.tex", "teagle_report.tex"])
def test_doc_has_a_decision_methodology_table(doc):
    """Each capability doc must carry the 'how each call is made' table (label tab:calls), so a reader sees
    the decision rule behind each class/superfamily/family/domain call, not only the vocabulary."""
    text = _doc(doc)
    if text is None:
        pytest.skip(f"{doc} not present")
    assert r"\label{tab:calls}" in text, f"{doc} is missing the decision-methodology table (tab:calls)"
    for rule in ("N-terminal", "C-terminal", "tyrosine recombinase", "80--80--80", "strongest-scoring"):
        assert rule in text, f"{doc} decision table omits the {rule!r} rule"


@pytest.mark.parametrize("doc", ["teagle_manual.tex", "teagle_report.tex"])
def test_doc_keeps_wicker_order_and_superfamily_ranks_distinct(doc):
    text = _doc(doc)
    if text is None:
        pytest.skip(f"{doc} not present")
    assert "Order & LINE (non-LTR)" in text
    assert "Order & DIRS group" in text
    assert "Order & LTR (superfamily undetermined)" in text
    assert "Superfamily & LINE (non-LTR)" not in text
    assert "Superfamily & DIRS group" not in text


def test_manual_distinguishes_withheld_ltr_origin_from_unresolved_ltr_superfamily():
    text = _doc("teagle_manual.tex")
    if text is None:
        pytest.skip("manual not present")
    flat = " ".join(text.split())
    assert "without RT or a trusted fragment architecture do not establish LTR origin" in flat
    assert "that result remains Class~I and is reported as \\texttt{LTR/unclassified}" in flat
    assert "the superfamily undetermined" in flat


def test_manual_states_tir_orientation_when_describing_identity():
    text = _doc("teagle_manual.tex")
    if text is None:
        pytest.skip("manual not present")
    flat = " ".join(text.split())
    assert "One TIR terminus is reverse-complemented before percentage identity is calculated" in flat
    assert "TIR termini are identical after reverse-complement alignment" not in flat
    assert "copies of an LTR or a TIR are identical at insertion" not in flat


def test_manual_herv_caption_matches_visible_headline_classification():
    text = _doc("teagle_manual.tex")
    if text is None:
        pytest.skip("manual not present")
    flat = " ".join(text.split())
    assert "The headline call is \\texttt{LTR/Gypsy}" in flat
    assert "the explanation identifies an env-bearing ERV lineage" in flat
    assert "classification banner reports an ERV" not in flat


def test_manual_does_not_treat_dfam_lookup_as_classifier_evidence():
    text = _doc("teagle_manual.tex")
    if text is None:
        pytest.skip("manual not present")
    flat = " ".join(text.split())
    assert "advisory database lookup" in flat
    assert "does not enter or alter the native domain/structure-based classification" in flat
    assert "independent line of evidence" not in flat


def test_manual_does_not_describe_calculated_primer_tm_as_measured():
    text = _doc("teagle_manual.tex")
    if text is None:
        pytest.skip("manual not present")
    assert "Primer3-calculated melting temperatures" in text
    assert "real melting temperatures" not in text


def test_docs_do_not_call_env_the_only_subgenomic_retroviral_transcript():
    texts = [
        _doc("teagle_manual.tex"),
        _doc("teagle_report.tex"),
        (Path(_ROOT) / "app" / "backend" / "teagle_core" / "retroviral.py").read_text(encoding="utf-8"),
        (Path(_ROOT) / "app" / "native" / "main.py").read_text(encoding="utf-8"),
        (Path(_ROOT) / "app" / "native" / "figures.py").read_text(encoding="utf-8"),
    ]
    assert all("env is the only subgenomic mRNA" not in text for text in texts if text is not None)
    assert all("gag-pro-pol is one continuous" not in text for text in texts if text is not None)
    assert all("gag--pro--pol} is one" not in text for text in texts if text is not None)
    assert all("gag–pro–pol intron (fused polyprotein)" not in text for text in texts if text is not None)
    assert all("exact bases cannot be read from the proviral DNA" not in text for text in texts if text is not None)
    assert all("exact splice bases cannot be recovered from proviral DNA" not in text for text in texts if text is not None)


def test_docs_do_not_present_repeat_identity_as_a_calibrated_age_limit():
    texts = [
        _doc("teagle_manual.tex"),
        _doc("teagle_report.tex"),
        (Path(_ROOT) / "app" / "native" / "main.py").read_text(encoding="utf-8"),
    ]
    assert all("age limit" not in text for text in texts if text is not None)


def test_docs_describe_primer_thermodynamics_as_model_predictions():
    manual = _doc("teagle_manual.tex")
    report = _doc("teagle_report.tex")
    if manual is None or report is None:
        pytest.skip("manual or report not present (report/ is a gitignored working dir)")
    manual = " ".join(manual.split())
    report = " ".join(report.split())
    assert "using real thermodynamics" not in manual
    assert "using real thermodynamics" not in report
    assert "nearest-neighbor thermodynamic models" in manual
    assert "nearest-neighbor thermodynamic models" in report


def test_installer_figure_caption_matches_component_row_layout():
    manual = _doc("teagle_manual.tex")
    if manual is None:
        pytest.skip("manual not present (report/ is a gitignored working dir)")
    assert "above its own button" not in manual
    assert "alongside its own button" in manual


def test_report_primer_summary_tracks_the_publication_table():
    """The report must retain the current in-silico primer scope and observed controlled errors."""
    text = _doc("teagle_report.tex")
    if text is None:
        pytest.skip("report not present")
    rows = _publication_table("table5_primer_and_reproducibility.csv")
    loci = [row for row in rows if row["Unit"].startswith("PRIMER")]
    controlled = next(row for row in rows if row["Unit"] == "Controlled scenarios")
    assert len(loci) == 3
    assert "Each of three de-novo loci" in text
    assert all(row["Candidates/tests"] == "5" for row in loci)
    assert controlled["Candidates/tests"] == "11"
    assert "eleven controlled off-target scenarios" in text
    assert "one false positive and one false negative" in text


@pytest.mark.parametrize("doc", ["teagle_manual.tex", "teagle_report.tex"])
def test_doc_lists_every_reported_domain_code_pfam_accession(doc):
    """Each Pfam accession in the bundled panel must appear in the doc's domain table — the table is the
    ground-truth enumeration the user asked for, so a model absent from it is a documentation defect."""
    text = _doc(doc)
    if text is None:
        pytest.skip(f"{doc} not present")
    for hmm, (code, label, cls, pfam) in domains.DOMAIN_INFO.items():
        assert pfam in text, f"{doc} omits {pfam} ({hmm}) from its domain-panel table"


def test_readme_version_badge_matches_the_source_version():
    from teagle_core import __version__
    stated = re.search(r"badge/version-([0-9.]+)-", _readme())
    assert stated, "README version badge missing"
    assert stated.group(1) == __version__


@pytest.mark.parametrize("relpath", ("report/teagle_manual.tex", "report/teagle_report.tex"))
def test_release_document_version_matches_the_source_version(relpath):
    """Ignored release documents must not retain an older version than the shipped build."""
    from teagle_core import __version__
    path = os.path.join(_ROOT, relpath)
    if not os.path.exists(path):
        pytest.skip(f"{relpath} not present in this checkout")
    stated = re.search(r"\\newcommand\{\\teagleversion\}\{([0-9.]+)\}", _read(relpath))
    assert stated, f"{relpath} does not define \\teagleversion"
    assert stated.group(1) == __version__


@pytest.mark.parametrize("doc", ["teagle_manual.tex", "teagle_report.tex"])
def test_genome_scan_docs_do_not_claim_a_fixed_runtime(doc):
    text = _doc(doc)
    if text is None:
        pytest.skip(f"{doc} not present")
    flat = " ".join(text.split())
    assert not re.search(r"(?:one|1)\s*(?:to|[-–])\s*2\s*(?:minutes?|mins?)", text, re.IGNORECASE)
    assert "typically scanned" not in text.lower()
    assert ("Runtime depends on genome size, the installed isPcr build and cached RefSeq assembly, local "
            "hardware, and the fixed candidate-product-size cap.") in flat
    assert "in-progress state while the scan runs and reports a timeout as a failure" in flat


def test_native_genome_scan_wording_matches_runtime_behavior():
    native = _read("app/native/main.py")
    wsl = _read("app/backend/teagle_core/wsl.py")
    flat = " ".join(native.split())
    assert not re.search(r"(?:one|1)\s*(?:to|[-–~])\s*2\s*(?:minutes?|mins?)", native, re.IGNORECASE)
    assert "fast once the genome is cached" not in native.lower()
    assert "future scans are fast" not in native.lower()
    assert ("Runtime depends on genome size, the installed isPcr build and cached RefSeq assembly, local "
            "hardware, and the fixed candidate-product-size cap") in flat
    assert "_render_genome_busy" in native
    assert "max_size: int = 4000" in wsl
    assert "timeout: int = 600" in wsl
    assert '"error": f"isPcr timed out after {timeout}s"' in wsl


def test_current_release_text_tracks_version_and_schema_contracts():
    """Current shipped text must distinguish 4.0 behavior from version-locked 3.9 evidence."""
    from teagle_core import __version__, provenance
    import engine

    citation = _read("CITATION.cff")
    stated = re.search(r"(?m)^version:\s*([0-9.]+)\s*$", citation)
    assert stated and stated.group(1) == __version__
    released = re.search(r"(?m)^date-released:\s*(\d{4}-\d{2}-\d{2})\s*$", citation)
    assert released, "a published release must carry its release date"

    readme = _read("README.md")
    app_readme = _read("app/README.md")
    changelog = _read("CHANGELOG.md")
    try_it = _read("docs/TRY-IT-YOURSELF.md")
    combined = "\n".join((readme, app_readme, changelog, try_it))
    assert f"## [{__version__}] — {released.group(1)}" in changelog
    assert f"Manifest schema {provenance.SCHEMA_VERSION}" in combined
    assert (f"result schema {engine.ANALYSIS_RESULT_SCHEMA}" in combined or
            f"analysis-result schema {engine.ANALYSIS_RESULT_SCHEMA}" in combined)
    assert f"source-context schema {provenance.SOURCE_CONTEXT_VERSION}" in combined
    assert "unkeyed integrity" in combined
    assert "clean-machine" in combined
    for required in (
        "paired terminal direct repeats", "core capsid/matrix GAG",
        "failed domain stage", "version-locked to TEagle 3.9.0",
    ):
        assert required in combined
    assert "Paired terminal direct repeats plus reverse transcriptase and an envelope-domain hit" in app_readme
    assert "Paired LTRs plus an envelope-domain hit" not in app_readme
    for stale in (
        "under release validation", "has not yet been published",
        "No 4.0.0 installer or manual asset is public yet",
        "still awaiting rendering and page-by-page release QA",
    ):
        assert stale not in combined


def test_current_docs_state_native_and_legacy_export_boundaries():
    readme = _read("README.md")
    app_readme = _read("app/README.md")
    for required in (
        "XLSX/CSV/TSV", "SVG/PNG/vector PDF", "GFF3/BED/FASTA",
        "provenance JSON", "complete-result JSON",
    ):
        assert required in readme
    assert "legacy browser" in readme.lower()
    assert "disables XLSX, PDF, provenance JSON, and complete-result JSON" in app_readme


def test_screenshot_index_records_completed_40_capture_contract():
    text = _read("docs/img/README.md")
    for required in (
        "native TEagle 4.0.0 app", "light theme at 1920x1080",
        "1500x1150 window and UI scale 1.0", "Every accepted figure is a real run",
        "No figure contains synthesised results", "recapture sidecar",
        "dimensions and SHA-256 hashes",
    ):
        assert required in text
    assert "must be recaptured" not in text


def test_manual_cites_subregion_picker_with_coordinate_contract():
    manual_path = os.path.join(_ROOT, "report", "teagle_manual.tex")
    if not os.path.exists(manual_path):
        pytest.skip("report/teagle_manual.tex not present in this checkout")
    manual = _read("report/teagle_manual.tex")
    assert "{subregion_picker.png}" in manual
    assert "1-based, inclusive interval" in manual
    assert "0-based, half-open span" in manual
    assert "primer design or splice detection" in manual


def test_manual_keeps_ltr_preposition_with_its_subject():
    manual = _read("report/teagle_manual.tex")

    def assert_phrase_kept(text):
        assert "For~LTR~retrotransposons" in text
        assert "For~LTR retrotransposons" not in text
        assert "For LTR retrotransposons" not in text

    assert_phrase_kept(manual)
    with pytest.raises(AssertionError):
        assert_phrase_kept(manual.replace("For~LTR~retrotransposons", "For~LTR retrotransposons", 1))


def test_manual_keeps_qpcr_gc_clamp_together():
    manual = _read("report/teagle_manual.tex")

    def assert_phrase_kept(text):
        assert "$3'$ GC~clamp." in text
        assert "$3'$ GC clamp." not in text

    assert_phrase_kept(manual)
    with pytest.raises(AssertionError):
        assert_phrase_kept(manual.replace("$3'$ GC~clamp.", "$3'$ GC clamp.", 1))


def test_manual_keeps_sensitivity_item_clear_of_the_footer():
    manual = _read("report/teagle_manual.tex")
    guard = "\\needspace{6\\baselineskip}\n\\item[Sensitivity]"

    def assert_guard(text):
        assert guard in text

    assert_guard(manual)
    with pytest.raises(AssertionError):
        assert_guard(manual.replace(guard, r"\item[Sensitivity]", 1))


def test_report_keeps_terminal_run_length_with_its_criterion():
    report = _read("report/teagle_report.tex")
    kept = r"length~$r \ge r_{\min}$"
    breakable = r"length $r \ge r_{\min}$"

    def assert_phrase_kept(text):
        assert kept in text
        assert breakable not in text

    assert_phrase_kept(report)
    with pytest.raises(AssertionError):
        assert_phrase_kept(report.replace(kept, breakable, 1))


def test_report_keeps_no_longer_together():
    report = _read("report/teagle_report.tex")
    kept = "no~longer"
    breakable = "no longer"

    def assert_phrase_kept(text):
        assert kept in text
        assert breakable not in text

    assert_phrase_kept(report)
    with pytest.raises(AssertionError):
        assert_phrase_kept(report.replace(kept, breakable, 1))


def test_report_keeps_stacked_rows_together():
    report = _read("report/teagle_report.tex")
    kept = "stacked~rows"
    breakable = "stacked rows"

    def assert_phrase_kept(text):
        assert kept in text
        assert breakable not in text

    assert_phrase_kept(report)
    with pytest.raises(AssertionError):
        assert_phrase_kept(report.replace(kept, breakable, 1))


def test_report_keeps_hue_is_together():
    report = _read("report/teagle_report.tex")
    kept = "Hue~is"
    breakable = "Hue is"

    def assert_phrase_kept(text):
        assert kept in text
        assert breakable not in text

    assert_phrase_kept(report)
    with pytest.raises(AssertionError):
        assert_phrase_kept(report.replace(kept, breakable, 1))


def test_manual_contents_card_navigation_has_no_gap_or_duplicate():
    manual = _read("report/teagle_manual.tex")
    navigation = re.compile(
        r"\\section\{Loading a sequence\}"
        r"|\\section\{Card (?P<section>\d{2}) --- (?P<section_title>[^{}]+)\}"
    )
    expected = [
        (1, "Loading a sequence"),
        (2, "Classification and structure"),
        (3, "Searching for a family match (optional)"),
        (4, "Designing PCR primers"),
        (5, "Predicting amplification (in-silico PCR)"),
        (6, "Whole-genome off-target scan"),
        (7, "The run provenance"),
    ]
    target = "\n".join((
        r"\section{Card 06 --- Whole-genome off-target scan}",
        r"\label{sec:genome-scan}",
    ))

    def assert_navigation(text):
        entries = []
        for match in navigation.finditer(text):
            if match.group(0) == r"\section{Loading a sequence}":
                entries.append((1, "Loading a sequence"))
            else:
                number = int(match.group("section"))
                title = match.group("section_title")
                entries.append((number, title))
        assert entries == expected
        assert target in text
        assert r"\addcontentsline{toc}" not in text

    card_06 = r"\section{Card 06 --- Whole-genome off-target scan}"
    assert_navigation(manual)
    with pytest.raises(AssertionError):
        assert_navigation(manual.replace(card_06 + "\n", "", 1))
    with pytest.raises(AssertionError):
        assert_navigation(manual.replace(card_06, card_06 + "\n" + card_06, 1))
    with pytest.raises(AssertionError):
        assert_navigation(manual + "\n" + r"\addcontentsline{toc}{section}{Card 06}")


def test_version_locked_benchmark_docs_match_their_derived_evidence_version():
    """Benchmark evidence keeps the version that generated it; an application patch must not relabel it."""
    with io.open(os.path.join(_ROOT, "reviewer_package", "results", "provenance_summary.json"),
                 encoding="utf-8") as handle:
        summary = json.load(handle)
    versions = summary["version_coverage"]["teagle_versions"]
    assert len(versions) == 1
    evidence_version, evidence_cases = next(iter(versions.items()))
    assert evidence_cases == summary["case_count"]
    assert f"TEagle {evidence_version}." in _read("reviewer_package/README.md")
    spec = _read("benchmarks/protocol/BENCHMARK_SPEC_v1.md")
    assert f"current execution pin is **{evidence_version}**" in spec
    # The current manuscript figure set is independent of the historical reviewer package.
    figure_index = os.path.join(_ROOT, "manuscript", "figures", "figure-index.md")
    if os.path.exists(figure_index):
        report = _read("report/teagle_report.tex")
        match = re.search(r"reported scientific results were generated with .*?~([0-9.]+),", report)
        assert match, "technical report must name the version that generated protocol 2.3.6 evidence"
        benchmark_version = match.group(1)
        current = _read("manuscript/figures/figure-index.md")
        assert f"TEagle {benchmark_version}" in current
        assert "protocol 2.3.6" in current


def test_internal_manuscript_holdout_summary_is_derived_from_current_tables():
    """Internal benchmark prose must follow the frozen protocol 2.3.6 presentation tables."""
    scope = next(
        row for row in _publication_table("table1_benchmark_scope.csv")
        if row["Panel"] == "External holdout"
    )
    rows = {row["Rank"].lower(): row for row in _publication_table("table2_selective_classification.csv")}
    text = _read("manuscript/manuscript.md")
    assert f"{scope['Rows']} independent sequences" in text
    assert rows["class"]["Correct/answered"] in text
    assert rows["order"]["Correct/answered"] in text
    assert rows["superfamily"]["Correct/answered"] in text


def test_manuscript_protocol_status_and_scope_are_explicit():
    text = _read("manuscript/manuscript.md")
    flat = " ".join(text.split())
    assert "protocol 2.3.6" in flat
    assert "10.5281/zenodo.22091013" in flat
    assert "identifies the preregistration package, not a results deposit" in flat
    assert "result-informed" in flat
    assert "DEV-0008" in flat
    assert "All 2.3.5 outputs were excluded" in flat
    assert "No wet-lab primer outcome was observed" in flat
    assert "not population estimates" in flat


def test_manuscript_primary_results_track_current_publication_tables():
    source = _read("manuscript/manuscript.md")
    classification = _publication_table("table2_selective_classification.csv")
    structure = _publication_table("table4_structure_and_domains.csv")[0]
    primers = _publication_table("table5_primer_and_reproducibility.csv")
    for row in classification:
        assert row["Answered"] in source
        assert row["Correct/answered"] in source
        assert row["Coverage"] in source
    assert f"{int(structure['N or truth']):,}" in source
    assert structure["Complete/TP"] in source
    assert structure["Incomplete/FN"] in source
    assert "98.6% precision, 29.7% operational recall and 45.7% F1" in source
    controlled = next(row for row in primers if row["Unit"] == "Controlled scenarios")
    reproducibility = next(row for row in primers if row["Unit"] == "Reproducibility gates")
    assert f"{controlled['Candidates/tests']} controlled off-target scenarios" in source
    assert reproducibility["Candidates/tests"] == "8"
    assert "All eight reproducibility checks passed" in source


def test_preprint_layout_source_requires_landscape_native_tables_and_caption_attachment():
    """Source-level guard for the DOCX layout policy; the artifact is checked below when present locally."""
    source = _read("manuscript/build_docx.py")
    for required in (
        "WD_ORIENT.LANDSCAPE",
        "WD_ORIENT.PORTRAIT",
        "w:tblHeader",
        "w:cantSplit",
        "keep_with_next",
        "keep_together",
        "WD_ALIGN_PARAGRAPH.LEFT",
        "WD_ALIGN_PARAGRAPH.JUSTIFY",
        "restart\"), \"continuous\"",
        "w:instrText",
        'instruction.text = " PAGE "',
        "_page_number_footer(section)",
    ):
        assert required in source


def test_manuscript_scientific_typography_source_is_explicit():
    """Scientific notation and biological names must carry unambiguous source markup."""
    text = _read("manuscript/manuscript.md")
    assert "<sup>1,2</sup>" in text and "<sup>2,3</sup>" in text
    assert "3′" in text and "5′" in text
    assert not re.search(r"[35]['’]", text)
    assert "*E*" in text
    assert "Pol polyprotein" in text and "pol polyprotein" not in text


def test_manuscript_claim_bearing_prose_uses_direct_scoped_wording():
    """Guard the reviewed claim passages without banning necessary scientific contrasts globally."""
    text = _read("manuscript/manuscript.md")
    flat = " ".join(text.split())
    for stale in (
        "Beneath this practical gap lies an epistemic one",
        "What is scarce is not the idea but its reach",
        "Two consequences follow",
        "More informative is that",
        "One property of this comparison has to be stated plainly",
        "The result that governs how the rest of this comparison must be read",
        "One qualification belongs with that statement",
        "we say so here rather than leave it to be inferred",
        "The decomposition isolates the observed cost",
        "The second observation concerns scope",
        "What the survey does support is narrower",
        "The general point survives whatever the count",
        "Both push in the same direction, making the recovery reported here optimistic",
        "The ground truth for the assay layer is external in a way a classification label is not",
        "[8, software documentation]",
        "DANTE_LTR reports each element",
        "The gap between the two is the honest measure",
        "informative rather than random",
        "A provenance claim is worth only as much as its weakest property",
        "Stability: does the same input yield the same seal",
        "Sensitivity: do different inputs yield different seals",
        "Coverage: does the seal record every threshold",
        "Classification is compared with literature-derived labels",
        "The idea of withholding rather than guessing is not new",
        "Withholding a call where the evidence cannot support one is not a new idea",
    ):
        assert stale not in flat
    for required in (
        "answered accuracy separately from coverage",
        "operational detection separately from conditional localization",
        "computational primer outcomes separately from wet-lab evidence",
        "Wilson intervals for answered accuracy are reported as case-level descriptive intervals",
        "organism-unit cluster bootstraps and block sign-flip tests were prespecified sensitivity analyses",
        "fixed-universe macro precision, recall and F1 were not estimable",
        "do not make the correction outcome-blind",
        "They do not support genome-scale annotation",
        "The primer results are computational, not biological validation",
    ):
        assert required in flat


def test_optional_local_manuscript_reports_all_three_reliability_axes():
    """The manuscript must not omit the HMMER call-confidence axis used by the shipped build."""
    _require_manuscript()
    flat = " ".join(_read("manuscript/manuscript.md").split())
    assert "Classification and its three reliability axes" in flat
    assert "per-domain call confidence" in flat
    assert "domain completeness" in flat
    assert "coding-structure" in flat
    assert "Classification and its two reliability axes" not in flat
    assert "Reliability is reported on two axes" not in flat


def test_docx_inline_renderer_uses_true_superscript_and_subscript_runs():
    """HTML-like scientific markup must become Word run properties, not smaller baseline text."""
    _require_manuscript()
    import sys as _sys
    from docx import Document

    manuscript_dir = os.path.join(_ROOT, "manuscript")
    if manuscript_dir not in _sys.path:
        _sys.path.insert(0, manuscript_dir)
    import build_docx

    paragraph = Document().add_paragraph()
    build_docx.add_inline(paragraph, "10<sup>−6</sup> and H<sub>2</sub>O")
    exponent = next(run for run in paragraph.runs if run.text == "−6")
    subscript = next(run for run in paragraph.runs if run.text == "2")
    assert exponent.font.superscript is True
    assert exponent.font.subscript is False
    assert subscript.font.subscript is True
    assert subscript.font.superscript is False


def test_optional_local_preprint_artifact_has_readable_table_and_figure_structure():
    """Optional local-artifact check: generated manuscript files are gitignored in some CI checkouts."""
    path = os.path.join(_ROOT, "manuscript", "manuscript_preprint.docx")
    if not os.path.exists(path):
        pytest.skip("preprint artifact not present (manuscript/ is a gitignored working directory)")
    from docx import Document
    from docx.oxml.ns import qn

    document = Document(path)
    landscape_sections = [section for section in document.sections
                          if section.page_width > section.page_height]
    portrait_sections = [section for section in document.sections
                         if section.page_width < section.page_height]
    assert len(document.tables) == 6
    assert len(landscape_sections) == len(document.tables)
    base_portrait = portrait_sections[0]
    base_margins = (
        base_portrait.left_margin, base_portrait.right_margin,
        base_portrait.top_margin, base_portrait.bottom_margin,
    )
    assert all((section.left_margin, section.right_margin,
                section.top_margin, section.bottom_margin) == base_margins
               for section in portrait_sections)
    for section in document.sections:
        numbering = section._sectPr.find(qn("w:lnNumType"))
        assert numbering is not None
        assert numbering.get(qn("w:restart")) == "continuous"
        footer_xml = section.footer._element.xml
        assert re.search(r"<w:instrText[^>]*>\s*PAGE\s*</w:instrText>", footer_xml)
        field_types = set(re.findall(r'w:fldCharType="([^"]+)"', footer_xml))
        assert {"begin", "separate", "end"}.issubset(field_types)
    for table in document.tables:
        assert table.rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader")) is not None
        for row in table.rows:
            assert row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None
    figure_paragraphs = [paragraph for paragraph in document.paragraphs
                         if "<w:drawing" in paragraph._p.xml]
    assert len(figure_paragraphs) == 6
    assert all(paragraph.paragraph_format.keep_with_next for paragraph in figure_paragraphs)
    captions = [paragraph for paragraph in document.paragraphs
                if re.match(r"^Figure \d+\.", paragraph.text)]
    assert len(captions) == 6
    assert all(paragraph.paragraph_format.keep_together for paragraph in captions)
    table_captions = [paragraph for paragraph in document.paragraphs
                      if re.match(r"^Table \d+\.", paragraph.text)]
    assert len(table_captions) == 6
    assert all(paragraph.paragraph_format.keep_together for paragraph in table_captions)
    assert all(paragraph.paragraph_format.keep_with_next for paragraph in table_captions)


def test_optional_local_manuscript_table_captions_do_not_split_across_pages():
    """Every manuscript rendition keeps each complete table caption in one paragraph box."""
    from docx import Document

    for filename in ("manuscript.docx", "manuscript_preprint.docx"):
        path = os.path.join(_ROOT, "manuscript", filename)
        if not os.path.exists(path):
            pytest.skip(f"{filename} not present")
        captions = [paragraph for paragraph in Document(path).paragraphs
                    if re.match(r"^Table \d+\.", paragraph.text)]
        assert len(captions) == 6
        assert all(paragraph.paragraph_format.keep_together for paragraph in captions)


def test_optional_local_manuscript_artifacts_preserve_scientific_run_properties():
    """Inspect the generated OOXML-level run styling for notation, taxa, genes and proteins."""
    from docx import Document

    for filename in ("manuscript.docx", "manuscript_preprint.docx"):
        path = os.path.join(_ROOT, "manuscript", filename)
        if not os.path.exists(path):
            pytest.skip(f"{filename} not present")
        document = Document(path)
        paragraphs = list(document.paragraphs)
        cells = [cell for table in document.tables for row in table.rows for cell in row.cells]
        paragraph_runs = [run for paragraph in paragraphs for run in paragraph.runs]
        runs = list(paragraph_runs)
        runs.extend(run for cell in cells for paragraph in cell.paragraphs for run in paragraph.runs)
        full_text = " ".join([paragraph.text for paragraph in paragraphs] + [cell.text for cell in cells])

        assert "<sup>" not in full_text and "<sub>" not in full_text
        assert "3′" in full_text and "5′" in full_text
        assert not re.search(r"[35]['’]", full_text)
        for exponent in ("1,2", "2,3", "2"):
            matching = [run for run in paragraph_runs if run.text == exponent]
            assert matching and all(run.font.superscript is True for run in matching)
        assert any(run.text == "E" and run.font.italic is True for run in runs)
        assert any("Pol polyprotein" in run.text and run.font.italic is False for run in runs)


def test_optional_local_manuscript_and_supplement_word_packages_use_only_times_new_roman():
    """Every Word text role, theme fallback, bullet and equation must resolve to Times New Roman."""
    paths = [
        os.path.join(_ROOT, "manuscript", "manuscript.docx"),
        os.path.join(_ROOT, "manuscript", "manuscript_preprint.docx"),
        os.path.join(_ROOT, "manuscript", "tables", "table4_tool_comparison.docx"),
        os.path.join(_ROOT, "manuscript", "tables", "tableS1_tool_comparison_sourced.docx"),
    ]
    present = [path for path in paths if os.path.exists(path)]
    if not present:
        pytest.skip("local manuscript artifacts are not present")
    for path in present:
        with zipfile.ZipFile(path) as package:
            word_xml = b"\n".join(
                package.read(name) for name in package.namelist()
                if name.startswith("word/") and name.endswith(".xml")
            )
        rfonts = re.findall(rb'<w:rFonts\b([^>]*)/?>', word_xml)
        fonts = {
            value
            for attributes in rfonts
            for value in re.findall(
                rb'w:(?:ascii|hAnsi|cs|eastAsia)="([^"]+)"', attributes
            )
        }
        theme_fonts = set(re.findall(rb'typeface="([^"]*)"', word_xml))
        math_fonts = set(re.findall(rb'<m:mathFont\b[^>]*m:val="([^"]+)"', word_xml))
        assert fonts == {b"Times New Roman"}, (os.path.basename(path), fonts)
        assert theme_fonts <= {b"Times New Roman"}, (os.path.basename(path), theme_fonts)
        assert math_fonts <= {b"Times New Roman"}, (os.path.basename(path), math_fonts)
        assert "\uf0b7".encode("utf-8") not in word_xml, os.path.basename(path)


def test_optional_local_manuscript_artifacts_fully_justify_body_prose():
    """Both Word variants justify body prose while leaving captions and references readable."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    for filename in ("manuscript.docx", "manuscript_preprint.docx"):
        path = os.path.join(_ROOT, "manuscript", filename)
        if not os.path.exists(path):
            pytest.skip(f"{filename} not present")
        document = Document(path)
        in_body = False
        in_references = False
        justified = 0
        references = 0
        for paragraph in document.paragraphs:
            if paragraph.text == "Abstract" and paragraph.style.name.startswith("Heading"):
                in_body = True
                continue
            if paragraph.text == "References" and paragraph.style.name.startswith("Heading"):
                in_references = True
                continue
            if not in_body or not paragraph.text.strip():
                continue
            is_heading = paragraph.style.name.startswith("Heading")
            is_list = paragraph.style.name == "List Bullet"
            is_caption = bool(re.match(r"^(Figure|Table) \d+\.", paragraph.text))
            is_figure = "<w:drawing" in paragraph._p.xml
            if is_heading or is_list or is_caption or is_figure:
                continue
            if in_references:
                assert paragraph.alignment == WD_ALIGN_PARAGRAPH.LEFT, (
                    filename, paragraph.text[:80]
                )
                references += 1
                continue
            assert paragraph.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY, (
                filename, paragraph.text[:80]
            )
            justified += 1
        assert justified > 60, (filename, justified)
        assert references == 25, (filename, references)
        title = next(p for p in document.paragraphs if p.text.startswith("TEagle: an evidence-traceable"))
        assert title.alignment != WD_ALIGN_PARAGRAPH.JUSTIFY
        captions = [p for p in document.paragraphs
                    if re.match(r"^(Figure|Table) \d+\.", p.text)]
        assert captions
        assert all(p.alignment == WD_ALIGN_PARAGRAPH.LEFT for p in captions)


def test_optional_local_manuscript_artifacts_are_double_spaced_throughout():
    """Mobile DNA asks for double spacing throughout the submitted manuscript."""
    from docx import Document

    for filename in ("manuscript.docx", "manuscript_preprint.docx"):
        path = os.path.join(_ROOT, "manuscript", filename)
        if not os.path.exists(path):
            pytest.skip(f"{filename} not present")
        document = Document(path)
        for paragraph in document.paragraphs:
            if not paragraph.text.strip():
                continue
            spacing = paragraph.paragraph_format.line_spacing
            if spacing is None:
                spacing = paragraph.style.paragraph_format.line_spacing
            assert spacing == 2.0, (filename, paragraph.text[:80], spacing)


def test_reviewer_corpus_scope_is_an_exact_mirror():
    """The reviewer copy must not silently diverge from the benchmark corpus-scope statement."""
    assert _read("reviewer_package/corpora/CORPUS_SCOPE.md") == _read("benchmarks/CORPUS_SCOPE.md")


def test_report_benchmark_spec_is_an_exact_mirror_when_maintained():
    """The optional, gitignored report mirror is exact when present in a local working tree."""
    report_spec = os.path.join(_ROOT, "report", "BENCHMARK_SPEC_v1.md")
    if not os.path.exists(report_spec):
        pytest.skip("report benchmark-spec mirror not present")
    assert _read("report/BENCHMARK_SPEC_v1.md") == _read("benchmarks/protocol/BENCHMARK_SPEC_v1.md")


def test_domains_tested_string_covers_every_emitted_domain_code():
    """DOMAINS_TESTED is the scope statement shown beside every 'not detected'. If a code can be emitted
    but is absent from that sentence, the honesty invariant is false."""
    from teagle_core import classify
    described = classify.DOMAINS_TESTED.lower()
    emitted = {info[0] for info in domains.DOMAIN_INFO.values()}
    expected_words = {"ORF1": "orf1", "EN": "endonuclease", "YR": "tyrosine recombinase",
                      "HEL": "helitron", "RT": "rt", "INT": "integrase", "ENV": "envelope",
                      "GAG": "gag", "PR": "protease", "RNaseH": "rnase h", "CHR": "chromodomain",
                      "TPase": "transposase"}
    for code in emitted:
        word = expected_words.get(code)
        assert word, f"domain code {code} has no expected wording — extend this test"
        assert word in described, f"{code} is emitted but not named in DOMAINS_TESTED"


# ---------------- curated assembly table ----------------
def test_every_curated_assembly_entry_is_well_formed():
    """Shape check, offline: a malformed accession or taxid would fail only when a user picked that
    organism and waited for a multi-GB download to fail."""
    import re
    from teagle_core import fetch
    for organism, meta in fetch.COORD_ASSEMBLIES.items():
        assert re.match(r"^GC[AF]_\d+\.\d+$", meta["assemblyAccession"]), (organism, meta)
        assert meta["taxid"].isdigit(), (organism, meta)
        assert meta["assemblyName"].strip(), organism


@pytest.mark.network
def test_every_curated_accession_resolves_to_the_organism_it_claims():
    """An accession that looks right and resolves to a different species is worse than none: it would
    silently scan the wrong genome. Verified against NCBI Datasets."""
    import json
    import urllib.request
    from teagle_core import fetch
    for organism, meta in fetch.COORD_ASSEMBLIES.items():
        acc = meta["assemblyAccession"]
        url = f"https://api.ncbi.nlm.nih.gov/datasets/v2/genome/accession/{acc}/dataset_report"
        report = (json.load(urllib.request.urlopen(url, timeout=60)).get("reports") or [{}])[0]
        got = report.get("organism", {})
        assert str(got.get("tax_id")) == meta["taxid"], (organism, acc, got.get("tax_id"))


def test_methods_panel_text_is_derived_from_the_profile_table():
    """The in-app methods disclosure once claimed a 21-model Pfam panel after it had grown to 30 — the app
    under-reporting its own method, in a UI string no docs test covered. It is now derived from
    domains.DOMAIN_INFO, so the count and the Pfam list cannot disagree with what the scan loads."""
    import os as _os, sys as _sys
    _os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    _native = _os.path.join(_os.path.dirname(_os.path.dirname(_os.path.abspath(__file__))), "app", "native")
    if _native not in _sys.path:
        _sys.path.insert(0, _native)
    import pytest as _pt
    _pt.importorskip("PySide6")
    from PySide6.QtWidgets import QApplication
    QApplication.instance() or QApplication([])
    import main
    from teagle_core import domains
    html = main.MainWindow._domain_panel_html()
    assert f"{len(domains.DOMAIN_INFO)} models" in html
    for _hmm, (_code, _label, _cls, pfam) in domains.DOMAIN_INFO.items():
        assert pfam in html, f"{pfam} missing from the methods panel text"


def test_tsd_method_text_states_the_window_the_detector_actually_searches():
    """Both TSD method strings claimed "4-12 bp" while find_tsd has searched from 2 bp since 3.2.1 — so the
    screen (and the exported Method column) told the user that the diagnostic 2 bp Tc1/Mariner TA, the exact
    case the detector was tuned to catch, could not be found. Both strings are now interpolated from
    structural.MIN_TSD/MAX_TSD; this fails if either drifts from the search bounds again."""
    import os as _os, sys as _sys
    _os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    _native = _os.path.join(_os.path.dirname(_os.path.dirname(_os.path.abspath(__file__))), "app", "native")
    if _native not in _sys.path:
        _sys.path.insert(0, _native)
    import pytest as _pt
    _pt.importorskip("PySide6")
    from PySide6.QtWidgets import QApplication
    QApplication.instance() or QApplication([])
    import inspect
    import main
    from teagle_core import structural

    # the constants must be what find_tsd actually defaults to, or deriving from them proves nothing
    sig = inspect.signature(structural.find_tsd)
    assert sig.parameters["min_tsd"].default == structural.MIN_TSD
    assert sig.parameters["max_tsd"].default == structural.MAX_TSD

    window = f"{structural.MIN_TSD}–{structural.MAX_TSD} bp"
    assert window in main.MainWindow._STRUCT_METHOD["TSD"], main.MainWindow._STRUCT_METHOD["TSD"]

    # _methods_html reads only self._src_html (citation links), so a stub `self` exercises the real string
    # without standing up a MainWindow.
    class _Stub:
        _domain_panel_html = staticmethod(main.MainWindow._domain_panel_html)

        def _src_html(self, _key):
            return ""

    html = main.MainWindow._methods_html(_Stub())
    assert window in html, "the Methods & databases disclosure states a stale TSD window"
    assert "Class / order / superfamily" in html
    assert "age limit" not in html
    assert "detection threshold" in html


def test_domains_tested_names_every_transposase_group_the_panel_can_call():
    """The scope sentence named five transposase groups while the panel also bundles Pfam's generic DDE
    family (PF03184, class dna:DDE) — which classify names outright as "DDE transposon". A user reading
    that line was told a family the classifier CAN call is not tested, which is the not-detected-vs-
    not-tested confusion the line exists to prevent. The phrase is now derived from DOMAIN_INFO; this
    fails if a transposase class is added to the panel and the sentence does not follow."""
    from teagle_core import classify, domains

    panel = {cls for code, _l, cls, _p in domains.DOMAIN_INFO.values()
             if code == "TPase" and cls.startswith("dna:")}
    assert panel, "no dna: transposase classes in the panel — update this test if that is intended"
    assert len(classify._tpase_groups()) == len(panel), (classify._tpase_groups(), panel)
    for group in classify._tpase_groups():
        assert group in classify.DOMAINS_TESTED, f"{group} missing from the tested-panel sentence"


def test_public_release_text_has_no_automated_author_attribution():
    forbidden = tuple("".join(parts).lower() for parts in (
        ("clau", "de"), ("chat", "gpt"), ("open", "ai"), ("co", "dex"),
    ))
    for relpath in ("README.md", "CHANGELOG.md", "app/README.md", "THIRD-PARTY-NOTICES.md"):
        path = os.path.join(_ROOT, relpath)
        if not os.path.exists(path):
            continue
        lower = _read(relpath).lower()
        assert not any(term in lower for term in forbidden), relpath


def test_mobile_dna_manuscript_submission_guardrails():
    """Keep the submitted prose aligned with the explicit Mobile DNA requirements and scope limits."""
    text = _read("manuscript/manuscript.md")
    flat = " ".join(text.split())
    lower = flat.lower()
    for heading in (
        "## Background", "## Implementation", "## Results", "## Discussion", "## Conclusions",
        "## Availability and requirements", "## Declarations", "## Abbreviations",
    ):
        assert heading in text
    for declaration in (
        "Ethics approval and consent to participate", "Consent for publication",
        "Availability of data and materials", "Competing interests", "Funding",
        "Authors' contributions", "Acknowledgements",
    ):
        assert declaration in text
    funding = (
        "Tuna Birgün received support from the Scientific and Technological Research Council of Türkiye "
        "(TÜBİTAK) through the 2210-A National MSc Scholarship Programme (Application No. 1649B022502808)."
    )
    assert text.count(funding) == 1
    assert "**Acknowledgements.** Not applicable." not in text
    assert "Tuna Birgün acknowledges the scholarship support stated under Funding." in text
    assert "This research received no separate project-specific grant." in text
    assert "GNU Affero General Public License v3.0 or later" in text
    assert text.count("Sibel Yılmaz") >= 2
    assert "Sibel Yilmaz" not in text
    for item in (
        "Project name", "Project home page", "Current stable release",
        "Archived protocol / persistent identifier", "Result archive", "Version benchmarked",
        "Operating system", "Programming language", "Other requirements", "License",
        "Any restrictions to use by non-academics",
    ):
        assert f"*{item}:*" in text
    assert "10.5281/zenodo.22091013" in text
    assert "This DOI identifies the preregistration, not the result archive." in text
    assert "executed result archive has not yet been assigned a public DOI" in text
    assert "10.5281/zenodo.21848814" not in text
    for forbidden in (
        "upper bound on performance", "proprietary", "benchmarks/", "benchmarks\\",
        "reviewer_package/", "reviewer_package\\", "STATISTICAL_ANALYSIS_PLAN.md",
        "repository test suite", "fails the build", "raw-output directory",
    ):
        assert forbidden not in text
    for unsupported_claim in (
        "general utility", "broad utility", "establishes population-level accuracy",
        "population-level accuracy is", "population-level accuracy was",
    ):
        assert unsupported_claim not in lower
    assert "integrated single-locus workflow" in lower
    assert "limited classification coverage" in lower
    assert "low structural recall" in lower
    assert "broad domain accuracy" in lower
    assert "wet-lab primer success" in lower
    assert "protocol 2.3.6 is transparently result-informed" in lower


def test_submission_assets_use_reader_facing_evidence_and_current_status():
    """Submission assets must not expose repository paths or repeat superseded review claims."""
    figure_sources = (
        "benchmarks/fig_workflow.py",
        "reviewer_package/scripts/fig_workflow.py",
        "manuscript/figures/figure-index.md",
        "manuscript/figures/fig_workflow.svg",
        "reviewer_package/figures/fig_workflow.svg",
    )
    for relpath in figure_sources:
        text = _read(relpath)
        assert "permissively licensed" not in text
    assert "Bundled — runs offline" in _read("manuscript/figures/fig_workflow.svg")

    table_paths = (
        "manuscript/tables/tableS1_tool_comparison_sourced.csv",
        "reviewer_package/results/tableS1_tool_comparison_sourced.csv",
    )
    forbidden_paths = ("app/backend/", "app/native/", "::_detector_parameters")
    for relpath in table_paths:
        text = _read(relpath)
        for forbidden in forbidden_paths:
            assert forbidden not in text
        assert "source and verification evidence supplied with the submission" in text

    for relpath in (
        "manuscript/tables/tableS1_tool_comparison_sourced.docx",
        "manuscript/tables/tableS1_tool_comparison_sourced.xlsx",
    ):
        with zipfile.ZipFile(os.path.join(_ROOT, relpath)) as package:
            payload = b"\n".join(package.read(name) for name in package.namelist())
        for forbidden in (b"app/backend/", b"app/native/", b"::_detector_parameters"):
            assert forbidden not in payload

    for relpath in (
        "manuscript/tables/table2_classification_accuracy.docx",
        "reviewer_package/results/table2_classification_accuracy.docx",
    ):
        with zipfile.ZipFile(os.path.join(_ROOT, relpath)) as package:
            payload = b"\n".join(package.read(name) for name in package.namelist())
        assert b"benchmarks/CORPUS_SCOPE.md" not in payload

    for relpath in (
        "manuscript/tables/table3_assay_validation.csv",
        "reviewer_package/results/table3_assay_validation.csv",
    ):
        header = _read(relpath).splitlines()[0]
        assert "Reference size (bp)" in header
        assert "Ground truth (bp)" not in header

    manuscript = " ".join(_read("manuscript/manuscript.md").split())
    assert "Clean-machine installation and uninstallation, and SmartScreen behaviour, were not tested" in manuscript
    assert "all required core dependencies are bundled" in manuscript
    assert "retrospectively replayed {{assay_cases}} published assays" in manuscript
    assert "without blinding" in manuscript
    assert "survive on a locked-down institutional laptop" not in manuscript
    assert "which capabilities are bundled and which require the optional WSL2 environment" in manuscript
    assert "Activation of tobacco retrotransposons during tissue culture" in manuscript
    assert "https://www.repeatmasker.org/ (accessed 7 August 2026" in manuscript

    reviewer = " ".join(_read("reviewer_package/README.md").split())
    assert "Phase 3 post-run" not in reviewer
    assert "Independent verification and submission readiness remain unclaimed" not in reviewer
    assert "benchmark evidence has persistent public access" in reviewer.lower()
